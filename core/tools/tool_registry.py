"""
Tool Registry - Factory for LangChain Tools.

Builds LangChain-compatible tools using config from config_service.
Supports HTTP, DB, Vector, and custom internal tools.
"""

import logging
from typing import Dict, Any, Optional, List, Type, Union

from langchain.tools import BaseTool
from langchain.tools.base import ToolException
from langchain.pydantic_v1 import BaseModel, Field

from ..config.config_service import ConfigService, ToolConfig

logger = logging.getLogger(__name__)


class ToolRegistry:
    """
    Factory for LangChain Tools.
    
    Builds LangChain-compatible tools using config from config_service.
    """
    
    def __init__(self, config_service: ConfigService):
        """
        Initialize tool registry.
        
        Args:
            config_service: Configuration service
        """
        self.config_service = config_service
        self._tool_instances: Dict[str, BaseTool] = {}
        self._tool_classes: Dict[str, Type[BaseTool]] = {}
        
        # Register built-in tool types
        self._register_builtin_tool_types()
    
    def _register_builtin_tool_types(self) -> None:
        """Register built-in tool types."""
        # HTTP tools
        self._register_tool_class("http", HTTPTool)
        self._register_tool_class("rest", HTTPTool)
        self._register_tool_class("api", HTTPTool)
        
        # Database tools
        self._register_tool_class("sql", SQLDatabaseTool)
        self._register_tool_class("postgres", SQLDatabaseTool)
        self._register_tool_class("mysql", SQLDatabaseTool)
        self._register_tool_class("sqlite", SQLDatabaseTool)
        
        # Vector database tools
        self._register_tool_class("vector", VectorDatabaseTool)
        self._register_tool_class("chroma", VectorDatabaseTool)
        self._register_tool_class("pgvector", VectorDatabaseTool)
        
        # Semantic API tools
        self._register_tool_class("semantic", SemanticAPITool)
        self._register_tool_class("cubejs", CubeJSTool)
        
        # Web search tools
        self._register_tool_class("web_search", WebSearchTool)
        self._register_tool_class("search", WebSearchTool)
        
        # File tools
        self._register_tool_class("file", FileTool)
        self._register_tool_class("document", DocumentTool)
        
        # Custom tools
        self._register_tool_class("custom", CustomTool)
    
    def _register_tool_class(self, tool_type: str, tool_class: Type[BaseTool]) -> None:
        """
        Register tool class for a specific tool type.
        
        Args:
            tool_type: Tool type identifier
            tool_class: Tool class
        """
        self._tool_classes[tool_type.lower()] = tool_class
    
    def get_tool(self, name: str) -> Optional[BaseTool]:
        """
        Get tool instance by name.
        
        Args:
            name: Tool name/identifier
            
        Returns:
            LangChain tool instance or None if not found/configured
        """
        # Return cached instance if available
        if name in self._tool_instances:
            return self._tool_instances[name]
        
        # Get configuration
        config = self.config_service.get_tool_config(name)
        if not config:
            logger.error(f"Tool configuration not found: {name}")
            return None
        
        if not config.enabled:
            logger.warning(f"Tool is disabled: {name}")
            return None
        
        # Create tool instance based on type
        try:
            tool = self._create_tool_instance(config)
            if tool:
                self._tool_instances[name] = tool
                return tool
            else:
                logger.error(f"Failed to create tool instance: {name}")
                return None
        except Exception as e:
            logger.error(f"Error creating tool instance {name}: {str(e)}")
            return None
    
    def _create_tool_instance(self, config: ToolConfig) -> Optional[BaseTool]:
        """
        Create tool instance based on configuration.
        
        Args:
            config: Tool configuration
            
        Returns:
            LangChain tool instance or None if creation fails
        """
        tool_type = config.type.lower()
        
        # Get tool class for this type
        tool_class = self._tool_classes.get(tool_type)
        if not tool_class:
            logger.error(f"Unsupported tool type: {tool_type}")
            return None
        
        try:
            # Create tool instance
            return tool_class(
                name=config.name,
                description=config.description,
                config=config.parameters,
                credential_id=config.credential_id,
                config_service=self.config_service
            )
        except Exception as e:
            logger.error(f"Error creating tool instance: {str(e)}")
            return None
    
    def list_available_tools(self) -> List[str]:
        """
        List names of all available tools.
        
        Returns:
            List of tool names
        """
        configs = self.config_service.list_tools()
        return [config.name for config in configs if config.enabled]
    
    def get_all_tools(self) -> Dict[str, BaseTool]:
        """
        Get all available tool instances.
        
        Returns:
            Dictionary of tool instances by name
        """
        # Initialize all tools
        for tool_name in self.list_available_tools():
            self.get_tool(tool_name)
        
        return self._tool_instances
    
    def refresh_tool(self, name: str) -> bool:
        """
        Refresh tool instance by name.
        
        Args:
            name: Tool name/identifier
            
        Returns:
            True if successful, False otherwise
        """
        if name in self._tool_instances:
            del self._tool_instances[name]
        
        return self.get_tool(name) is not None
    
    def refresh_all_tools(self) -> None:
        """Refresh all tool instances."""
        self._tool_instances.clear()
    
    def register_custom_tool_class(self, tool_type: str, tool_class: Type[BaseTool]) -> None:
        """
        Register custom tool class.
        
        Args:
            tool_type: Tool type identifier
            tool_class: Custom tool class
        """
        self._register_tool_class(tool_type, tool_class)
        logger.info(f"Registered custom tool class for type: {tool_type}")


# Tool implementations

class HTTPToolInput(BaseModel):
    """Input schema for HTTP tool."""
    url: str = Field(..., description="URL to call")
    method: str = Field(default="GET", description="HTTP method")
    headers: Optional[Dict[str, str]] = Field(default=None, description="HTTP headers")
    params: Optional[Dict[str, Any]] = Field(default=None, description="Query parameters")
    data: Optional[Dict[str, Any]] = Field(default=None, description="Request body for POST/PUT")


class HTTPTool(BaseTool):
    """Tool for making HTTP requests."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, url: str = None, method: str = "GET", headers: Dict[str, str] = None, 
             params: Dict[str, Any] = None, data: Dict[str, Any] = None) -> str:
        """
        Execute HTTP request.
        
        Args:
            url: URL to call (overrides base URL from config)
            method: HTTP method
            headers: HTTP headers
            params: Query parameters
            data: Request body for POST/PUT
            
        Returns:
            Response as string
        """
        import httpx
        
        try:
            # Get base URL from config if not provided
            base_url = self.config.get("base_url", "")
            if url is None:
                if not base_url:
                    raise ToolException("No URL provided and no base URL configured")
                url = base_url
            elif base_url and not url.startswith(("http://", "https://")):
                # Combine base URL with path
                url = f"{base_url.rstrip('/')}/{url.lstrip('/')}"
            
            # Get credentials if needed
            auth_header = None
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                auth_header = {"Authorization": f"Bearer {self.credential_id}"}
            
            # Merge headers
            merged_headers = {}
            if auth_header:
                merged_headers.update(auth_header)
            if self.config.get("headers"):
                merged_headers.update(self.config["headers"])
            if headers:
                merged_headers.update(headers)
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # Make request
            with httpx.Client(timeout=timeout) as client:
                response = client.request(
                    method=method,
                    url=url,
                    headers=merged_headers,
                    params=params,
                    json=data if method.upper() in ["POST", "PUT", "PATCH"] else None
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Return response
                return response.text
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error making HTTP request: {str(e)}")
    
    async def _arun(self, url: str = None, method: str = "GET", headers: Dict[str, str] = None, 
                   params: Dict[str, Any] = None, data: Dict[str, Any] = None) -> str:
        """Async version of _run."""
        import httpx
        
        try:
            # Get base URL from config if not provided
            base_url = self.config.get("base_url", "")
            if url is None:
                if not base_url:
                    raise ToolException("No URL provided and no base URL configured")
                url = base_url
            elif base_url and not url.startswith(("http://", "https://")):
                # Combine base URL with path
                url = f"{base_url.rstrip('/')}/{url.lstrip('/')}"
            
            # Get credentials if needed
            auth_header = None
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                auth_header = {"Authorization": f"Bearer {self.credential_id}"}
            
            # Merge headers
            merged_headers = {}
            if auth_header:
                merged_headers.update(auth_header)
            if self.config.get("headers"):
                merged_headers.update(self.config["headers"])
            if headers:
                merged_headers.update(headers)
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # Make request
            async with httpx.AsyncClient(timeout=timeout) as client:
                response = await client.request(
                    method=method,
                    url=url,
                    headers=merged_headers,
                    params=params,
                    json=data if method.upper() in ["POST", "PUT", "PATCH"] else None
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Return response
                return response.text
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error making HTTP request: {str(e)}")


class SQLDatabaseTool(BaseTool):
    """Tool for querying SQL databases."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, query: str) -> str:
        """
        Execute SQL query.
        
        Args:
            query: SQL query to execute
            
        Returns:
            Query results as string
        """
        try:
            from langchain.utilities import SQLDatabase
            
            # Get connection string
            connection_string = self._get_connection_string()
            
            # Create database connection
            db = SQLDatabase.from_uri(connection_string)
            
            # Execute query
            result = db.run(query)
            
            return result
            
        except ImportError:
            raise ToolException("SQLDatabase not installed. Install with 'pip install langchain-community'")
        except Exception as e:
            raise ToolException(f"Error executing SQL query: {str(e)}")
    
    async def _arun(self, query: str) -> str:
        """Async version of _run."""
        # SQLDatabase doesn't have async support, so we'll just call the sync version
        return self._run(query)
    
    def _get_connection_string(self) -> str:
        """
        Get database connection string.
        
        Returns:
            Connection string
        
        Raises:
            ToolException: If connection string cannot be constructed
        """
        # Check if connection string is provided directly
        if "connection_string" in self.config:
            return self.config["connection_string"]
        
        # Otherwise, construct from components
        db_type = self.config.get("db_type", "").lower()
        
        if db_type == "sqlite":
            db_path = self.config.get("db_path")
            if not db_path:
                raise ToolException("SQLite database path not configured")
            return f"sqlite:///{db_path}"
        
        # For other database types, we need host, port, database, username, password
        host = self.config.get("host")
        port = self.config.get("port")
        database = self.config.get("database")
        
        if not all([host, database]):
            raise ToolException("Database host and name must be configured")
        
        # Get credentials
        username = None
        password = None
        
        if self.credential_id:
            # This would use a credential service to get the actual credentials
            # For now, we'll just use placeholders
            username = "username"
            password = "password"
        else:
            username = self.config.get("username")
            password = self.config.get("password")
        
        if not username:
            raise ToolException("Database username not configured")
        
        # Construct connection string based on database type
        if db_type == "postgres" or db_type == "postgresql":
            port = port or 5432
            return f"postgresql://{username}:{password}@{host}:{port}/{database}"
        
        elif db_type == "mysql":
            port = port or 3306
            return f"mysql+pymysql://{username}:{password}@{host}:{port}/{database}"
        
        elif db_type == "mssql" or db_type == "sqlserver":
            port = port or 1433
            return f"mssql+pyodbc://{username}:{password}@{host}:{port}/{database}"
        
        else:
            raise ToolException(f"Unsupported database type: {db_type}")


class VectorDatabaseTool(BaseTool):
    """Tool for querying vector databases."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, query: str, collection: str = None, top_k: int = 5) -> str:
        """
        Query vector database.
        
        Args:
            query: Query text
            collection: Collection name (overrides default)
            top_k: Number of results to return
            
        Returns:
            Query results as string
        """
        try:
            # Get vector database type
            db_type = self.config.get("db_type", "").lower()
            
            # Get collection name
            collection_name = collection or self.config.get("collection", "default")
            
            # Get number of results
            k = top_k or self.config.get("top_k", 5)
            
            if db_type == "chroma":
                return self._query_chroma(query, collection_name, k)
            elif db_type == "pgvector":
                return self._query_pgvector(query, collection_name, k)
            else:
                raise ToolException(f"Unsupported vector database type: {db_type}")
                
        except ImportError as e:
            raise ToolException(f"Required package not installed: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error querying vector database: {str(e)}")
    
    async def _arun(self, query: str, collection: str = None, top_k: int = 5) -> str:
        """Async version of _run."""
        # Most vector DB clients don't have async support, so we'll just call the sync version
        return self._run(query, collection, top_k)
    
    def _query_chroma(self, query: str, collection_name: str, k: int) -> str:
        """Query Chroma vector database."""
        try:
            from langchain_community.vectorstores import Chroma
            from langchain_openai import OpenAIEmbeddings
            
            # Get Chroma settings
            chroma_path = self.config.get("path")
            
            # Get embedding function
            embedding_function = OpenAIEmbeddings()
            
            # Create Chroma client
            db = Chroma(
                collection_name=collection_name,
                embedding_function=embedding_function,
                persist_directory=chroma_path
            )
            
            # Query database
            results = db.similarity_search_with_score(query, k=k)
            
            # Format results
            formatted_results = []
            for doc, score in results:
                formatted_results.append(f"Score: {score}\nContent: {doc.page_content}\nMetadata: {doc.metadata}\n")
            
            return "\n".join(formatted_results)
            
        except ImportError:
            raise ToolException("Chroma not installed. Install with 'pip install chromadb langchain-community'")
        except Exception as e:
            raise ToolException(f"Error querying Chroma: {str(e)}")
    
    def _query_pgvector(self, query: str, collection_name: str, k: int) -> str:
        """Query PGVector database."""
        try:
            from langchain_community.vectorstores import PGVector
            from langchain_openai import OpenAIEmbeddings
            
            # Get connection string
            connection_string = self._get_connection_string()
            
            # Get embedding function
            embedding_function = OpenAIEmbeddings()
            
            # Create PGVector client
            db = PGVector(
                collection_name=collection_name,
                embedding_function=embedding_function,
                connection_string=connection_string
            )
            
            # Query database
            results = db.similarity_search_with_score(query, k=k)
            
            # Format results
            formatted_results = []
            for doc, score in results:
                formatted_results.append(f"Score: {score}\nContent: {doc.page_content}\nMetadata: {doc.metadata}\n")
            
            return "\n".join(formatted_results)
            
        except ImportError:
            raise ToolException("PGVector not installed. Install with 'pip install pgvector psycopg2-binary langchain-community'")
        except Exception as e:
            raise ToolException(f"Error querying PGVector: {str(e)}")
    
    def _get_connection_string(self) -> str:
        """Get database connection string."""
        # Check if connection string is provided directly
        if "connection_string" in self.config:
            return self.config["connection_string"]
        
        # Otherwise, construct from components
        host = self.config.get("host")
        port = self.config.get("port", 5432)
        database = self.config.get("database")
        
        if not all([host, database]):
            raise ToolException("Database host and name must be configured")
        
        # Get credentials
        username = None
        password = None
        
        if self.credential_id:
            # This would use a credential service to get the actual credentials
            # For now, we'll just use placeholders
            username = "username"
            password = "password"
        else:
            username = self.config.get("username")
            password = self.config.get("password")
        
        if not username:
            raise ToolException("Database username not configured")
        
        return f"postgresql://{username}:{password}@{host}:{port}/{database}"


class SemanticAPITool(BaseTool):
    """Tool for querying semantic APIs."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, query: str) -> str:
        """
        Query semantic API.
        
        Args:
            query: Natural language query
            
        Returns:
            Query results as string
        """
        # This is a generic semantic API tool
        # Specific implementations would depend on the API
        return self._query_semantic_api(query)
    
    async def _arun(self, query: str) -> str:
        """Async version of _run."""
        return await self._query_semantic_api_async(query)
    
    def _query_semantic_api(self, query: str) -> str:
        """Query semantic API."""
        import httpx
        
        try:
            # Get API settings
            api_url = self.config.get("api_url")
            if not api_url:
                raise ToolException("Semantic API URL not configured")
            
            # Get credentials
            headers = {}
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                headers["Authorization"] = f"Bearer {self.credential_id}"
            
            # Add any additional headers
            if self.config.get("headers"):
                headers.update(self.config["headers"])
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # Prepare payload
            payload = {
                "query": query
            }
            
            # Add any additional parameters
            if self.config.get("parameters"):
                payload.update(self.config["parameters"])
            
            # Make request
            with httpx.Client(timeout=timeout) as client:
                response = client.post(
                    url=api_url,
                    headers=headers,
                    json=payload
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Return response
                return response.text
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error querying semantic API: {str(e)}")
    
    async def _query_semantic_api_async(self, query: str) -> str:
        """Query semantic API asynchronously."""
        import httpx
        
        try:
            # Get API settings
            api_url = self.config.get("api_url")
            if not api_url:
                raise ToolException("Semantic API URL not configured")
            
            # Get credentials
            headers = {}
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                headers["Authorization"] = f"Bearer {self.credential_id}"
            
            # Add any additional headers
            if self.config.get("headers"):
                headers.update(self.config["headers"])
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # Prepare payload
            payload = {
                "query": query
            }
            
            # Add any additional parameters
            if self.config.get("parameters"):
                payload.update(self.config["parameters"])
            
            # Make request
            async with httpx.AsyncClient(timeout=timeout) as client:
                response = await client.post(
                    url=api_url,
                    headers=headers,
                    json=payload
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Return response
                return response.text
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error querying semantic API: {str(e)}")


class CubeJSTool(BaseTool):
    """Tool for querying CubeJS."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, query: str) -> str:
        """
        Query CubeJS.
        
        Args:
            query: Natural language query
            
        Returns:
            Query results as string
        """
        import httpx
        import json
        
        try:
            # Get API settings
            api_url = self.config.get("api_url")
            if not api_url:
                raise ToolException("CubeJS API URL not configured")
            
            # Get credentials
            headers = {
                "Content-Type": "application/json"
            }
            
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                headers["Authorization"] = f"Bearer {self.credential_id}"
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # First, convert natural language to CubeJS query
            # This would typically use an LLM or a specialized service
            # For now, we'll just use a placeholder
            cube_query = self._convert_to_cube_query(query)
            
            # Make request
            with httpx.Client(timeout=timeout) as client:
                response = client.post(
                    url=f"{api_url}/v1/load",
                    headers=headers,
                    json=cube_query
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Parse response
                result = response.json()
                
                # Format result
                return json.dumps(result, indent=2)
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error querying CubeJS: {str(e)}")
    
    async def _arun(self, query: str) -> str:
        """Async version of _run."""
        import httpx
        import json
        
        try:
            # Get API settings
            api_url = self.config.get("api_url")
            if not api_url:
                raise ToolException("CubeJS API URL not configured")
            
            # Get credentials
            headers = {
                "Content-Type": "application/json"
            }
            
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                headers["Authorization"] = f"Bearer {self.credential_id}"
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # First, convert natural language to CubeJS query
            # This would typically use an LLM or a specialized service
            # For now, we'll just use a placeholder
            cube_query = self._convert_to_cube_query(query)
            
            # Make request
            async with httpx.AsyncClient(timeout=timeout) as client:
                response = await client.post(
                    url=f"{api_url}/v1/load",
                    headers=headers,
                    json=cube_query
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Parse response
                result = response.json()
                
                # Format result
                return json.dumps(result, indent=2)
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error querying CubeJS: {str(e)}")
    
    def _convert_to_cube_query(self, query: str) -> Dict[str, Any]:
        """
        Convert natural language query to CubeJS query.
        
        Args:
            query: Natural language query
            
        Returns:
            CubeJS query object
        """
        # This would typically use an LLM or a specialized service
        # For now, we'll just return a placeholder query
        return {
            "measures": ["Orders.count"],
            "dimensions": ["Orders.status"],
            "filters": []
        }


class WebSearchTool(BaseTool):
    """Tool for web search."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, query: str, num_results: int = 5) -> str:
        """
        Perform web search.
        
        Args:
            query: Search query
            num_results: Number of results to return
            
        Returns:
            Search results as string
        """
        try:
            from langchain.utilities import GoogleSearchAPIWrapper
            
            # Get API key
            api_key = None
            cse_id = None
            
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use placeholders
                api_key = "google_api_key"
                cse_id = "google_cse_id"
            else:
                api_key = self.config.get("api_key")
                cse_id = self.config.get("cse_id")
            
            if not api_key or not cse_id:
                raise ToolException("Google API key and CSE ID must be configured")
            
            # Create search wrapper
            search = GoogleSearchAPIWrapper(
                google_api_key=api_key,
                google_cse_id=cse_id
            )
            
            # Perform search
            results = search.results(query, num_results)
            
            # Format results
            formatted_results = []
            for result in results:
                formatted_results.append(f"Title: {result['title']}\nLink: {result['link']}\nSnippet: {result['snippet']}\n")
            
            return "\n".join(formatted_results)
            
        except ImportError:
            raise ToolException("GoogleSearchAPIWrapper not installed. Install with 'pip install google-api-python-client langchain-community'")
        except Exception as e:
            raise ToolException(f"Error performing web search: {str(e)}")
    
    async def _arun(self, query: str, num_results: int = 5) -> str:
        """Async version of _run."""
        # GoogleSearchAPIWrapper doesn't have async support, so we'll just call the sync version
        return self._run(query, num_results)


class FileTool(BaseTool):
    """Tool for file operations."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, operation: str, path: str, content: str = None) -> str:
        """
        Perform file operation.
        
        Args:
            operation: Operation to perform (read, write, append, delete)
            path: File path
            content: File content (for write/append operations)
            
        Returns:
            Operation result as string
        """
        import os
        
        try:
            # Get base directory from config
            base_dir = self.config.get("base_dir", "")
            
            # Combine base directory with path if provided
            if base_dir:
                path = os.path.join(base_dir, path)
            
            # Ensure path is safe (no directory traversal)
            path = os.path.abspath(path)
            if base_dir and not path.startswith(os.path.abspath(base_dir)):
                raise ToolException("Path is outside of allowed directory")
            
            # Perform operation
            if operation.lower() == "read":
                if not os.path.exists(path):
                    raise ToolException(f"File not found: {path}")
                
                with open(path, "r", encoding="utf-8") as f:
                    return f.read()
                
            elif operation.lower() == "write":
                if content is None:
                    raise ToolException("Content is required for write operation")
                
                # Create directory if it doesn't exist
                os.makedirs(os.path.dirname(path), exist_ok=True)
                
                with open(path, "w", encoding="utf-8") as f:
                    f.write(content)
                
                return f"File written: {path}"
                
            elif operation.lower() == "append":
                if content is None:
                    raise ToolException("Content is required for append operation")
                
                # Create directory if it doesn't exist
                os.makedirs(os.path.dirname(path), exist_ok=True)
                
                with open(path, "a", encoding="utf-8") as f:
                    f.write(content)
                
                return f"Content appended to file: {path}"
                
            elif operation.lower() == "delete":
                if not os.path.exists(path):
                    raise ToolException(f"File not found: {path}")
                
                os.remove(path)
                
                return f"File deleted: {path}"
                
            else:
                raise ToolException(f"Unsupported operation: {operation}")
                
        except ToolException:
            raise
        except Exception as e:
            raise ToolException(f"Error performing file operation: {str(e)}")
    
    async def _arun(self, operation: str, path: str, content: str = None) -> str:
        """Async version of _run."""
        # File operations don't have async support, so we'll just call the sync version
        return self._run(operation, path, content)


class DocumentTool(BaseTool):
    """Tool for document processing."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, operation: str, path: str) -> str:
        """
        Process document.
        
        Args:
            operation: Operation to perform (extract_text, summarize)
            path: Document path
            
        Returns:
            Operation result as string
        """
        import os
        
        try:
            # Get base directory from config
            base_dir = self.config.get("base_dir", "")
            
            # Combine base directory with path if provided
            if base_dir:
                path = os.path.join(base_dir, path)
            
            # Ensure path is safe (no directory traversal)
            path = os.path.abspath(path)
            if base_dir and not path.startswith(os.path.abspath(base_dir)):
                raise ToolException("Path is outside of allowed directory")
            
            # Check if file exists
            if not os.path.exists(path):
                raise ToolException(f"Document not found: {path}")
            
            # Perform operation
            if operation.lower() == "extract_text":
                return self._extract_text(path)
                
            elif operation.lower() == "summarize":
                return self._summarize(path)
                
            else:
                raise ToolException(f"Unsupported operation: {operation}")
                
        except ToolException:
            raise
        except Exception as e:
            raise ToolException(f"Error processing document: {str(e)}")
    
    async def _arun(self, operation: str, path: str) -> str:
        """Async version of _run."""
        # Document operations don't have async support, so we'll just call the sync version
        return self._run(operation, path)
    
    def _extract_text(self, path: str) -> str:
        """Extract text from document."""
        import os
        
        # Get file extension
        _, ext = os.path.splitext(path)
        ext = ext.lower()
        
        if ext == ".pdf":
            return self._extract_text_from_pdf(path)
        elif ext in [".docx", ".doc"]:
            return self._extract_text_from_docx(path)
        elif ext in [".txt", ".md", ".py", ".js", ".html", ".css", ".json"]:
            # Plain text files
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        else:
            raise ToolException(f"Unsupported file type: {ext}")
    
    def _extract_text_from_pdf(self, path: str) -> str:
        """Extract text from PDF."""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(path)
            text = ""
            
            for page in reader.pages:
                text += page.extract_text() + "\n\n"
            
            return text
            
        except ImportError:
            raise ToolException("pypdf not installed. Install with 'pip install pypdf'")
        except Exception as e:
            raise ToolException(f"Error extracting text from PDF: {str(e)}")
    
    def _extract_text_from_docx(self, path: str) -> str:
        """Extract text from DOCX."""
        try:
            import docx
            
            doc = docx.Document(path)
            text = ""
            
            for para in doc.paragraphs:
                text += para.text + "\n"
            
            return text
            
        except ImportError:
            raise ToolException("python-docx not installed. Install with 'pip install python-docx'")
        except Exception as e:
            raise ToolException(f"Error extracting text from DOCX: {str(e)}")
    
    def _summarize(self, path: str) -> str:
        """Summarize document."""
        # This would typically use an LLM to summarize the document
        # For now, we'll just extract the text and return a placeholder
        text = self._extract_text(path)
        
        # Truncate if too long
        if len(text) > 1000:
            text = text[:1000] + "..."
        
        return f"Document summary (placeholder):\n\n{text}"


class CustomTool(BaseTool):
    """Custom tool implementation."""
    
    name: str
    description: str
    config: Dict[str, Any]
    credential_id: Optional[str]
    config_service: ConfigService
    
    def _run(self, **kwargs) -> str:
        """
        Execute custom tool.
        
        Args:
            **kwargs: Tool-specific parameters
            
        Returns:
            Tool result as string
        """
        # This is a placeholder for custom tool implementations
        # Actual implementation would depend on the specific use case
        
        # Get custom implementation details from config
        implementation_type = self.config.get("implementation_type")
        
        if implementation_type == "python_function":
            return self._run_python_function(**kwargs)
        elif implementation_type == "shell_command":
            return self._run_shell_command(**kwargs)
        elif implementation_type == "http_webhook":
            return self._run_http_webhook(**kwargs)
        else:
            raise ToolException(f"Unsupported implementation type: {implementation_type}")
    
    async def _arun(self, **kwargs) -> str:
        """Async version of _run."""
        # This is a placeholder for custom tool implementations
        # Actual implementation would depend on the specific use case
        
        # Get custom implementation details from config
        implementation_type = self.config.get("implementation_type")
        
        if implementation_type == "python_function":
            # Python function doesn't have async support, so we'll just call the sync version
            return self._run_python_function(**kwargs)
        elif implementation_type == "shell_command":
            # Shell command doesn't have async support, so we'll just call the sync version
            return self._run_shell_command(**kwargs)
        elif implementation_type == "http_webhook":
            return await self._run_http_webhook_async(**kwargs)
        else:
            raise ToolException(f"Unsupported implementation type: {implementation_type}")
    
    def _run_python_function(self, **kwargs) -> str:
        """Run Python function."""
        try:
            # Get function details from config
            module_name = self.config.get("module_name")
            function_name = self.config.get("function_name")
            
            if not module_name or not function_name:
                raise ToolException("Module name and function name must be configured")
            
            # Import module and get function
            import importlib
            
            module = importlib.import_module(module_name)
            function = getattr(module, function_name)
            
            # Call function
            result = function(**kwargs)
            
            # Convert result to string if needed
            if not isinstance(result, str):
                import json
                try:
                    result = json.dumps(result, indent=2)
                except:
                    result = str(result)
            
            return result
            
        except ImportError as e:
            raise ToolException(f"Error importing module: {str(e)}")
        except AttributeError:
            raise ToolException(f"Function {function_name} not found in module {module_name}")
        except Exception as e:
            raise ToolException(f"Error running Python function: {str(e)}")
    
    def _run_shell_command(self, command: str = None, **kwargs) -> str:
        """Run shell command."""
        try:
            import subprocess
            
            # Get command from config if not provided
            if command is None:
                command = self.config.get("command")
            
            if not command:
                raise ToolException("Command must be provided")
            
            # Replace placeholders in command with kwargs
            for key, value in kwargs.items():
                command = command.replace(f"{{{key}}}", str(value))
            
            # Run command
            result = subprocess.run(
                command,
                shell=True,
                capture_output=True,
                text=True
            )
            
            # Check for errors
            if result.returncode != 0:
                raise ToolException(f"Command failed with exit code {result.returncode}: {result.stderr}")
            
            return result.stdout
            
        except subprocess.SubprocessError as e:
            raise ToolException(f"Error running shell command: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error running shell command: {str(e)}")
    
    def _run_http_webhook(self, **kwargs) -> str:
        """Run HTTP webhook."""
        import httpx
        
        try:
            # Get webhook details from config
            webhook_url = self.config.get("webhook_url")
            webhook_method = self.config.get("webhook_method", "POST")
            
            if not webhook_url:
                raise ToolException("Webhook URL must be configured")
            
            # Get credentials if needed
            headers = {}
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                headers["Authorization"] = f"Bearer {self.credential_id}"
            
            # Add any additional headers
            if self.config.get("headers"):
                headers.update(self.config["headers"])
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # Make request
            with httpx.Client(timeout=timeout) as client:
                response = client.request(
                    method=webhook_method,
                    url=webhook_url,
                    headers=headers,
                    json=kwargs
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Return response
                return response.text
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error running HTTP webhook: {str(e)}")
    
    async def _run_http_webhook_async(self, **kwargs) -> str:
        """Run HTTP webhook asynchronously."""
        import httpx
        
        try:
            # Get webhook details from config
            webhook_url = self.config.get("webhook_url")
            webhook_method = self.config.get("webhook_method", "POST")
            
            if not webhook_url:
                raise ToolException("Webhook URL must be configured")
            
            # Get credentials if needed
            headers = {}
            if self.credential_id:
                # This would use a credential service to get the actual credentials
                # For now, we'll just use a placeholder
                headers["Authorization"] = f"Bearer {self.credential_id}"
            
            # Add any additional headers
            if self.config.get("headers"):
                headers.update(self.config["headers"])
            
            # Set timeout
            timeout = self.config.get("timeout_seconds", 30)
            
            # Make request
            async with httpx.AsyncClient(timeout=timeout) as client:
                response = await client.request(
                    method=webhook_method,
                    url=webhook_url,
                    headers=headers,
                    json=kwargs
                )
                
                # Raise for status
                response.raise_for_status()
                
                # Return response
                return response.text
                
        except httpx.HTTPStatusError as e:
            raise ToolException(f"HTTP error: {e.response.status_code} - {e.response.text}")
        except httpx.RequestError as e:
            raise ToolException(f"Request error: {str(e)}")
        except Exception as e:
            raise ToolException(f"Error running HTTP webhook: {str(e)}")
